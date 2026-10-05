"""
Renders the portfolio dashboard and risk log as a real, multi-sheet
.xlsx workbook via openpyxl -- an actual file a stakeholder could open,
not a printed table.
"""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from src.analysis import flag_attention_needed, portfolio_dashboard
from src.data_sources import MILESTONES, PROJECTS, RISK_LOG, STATUS_REPORTS

_HEADER_FILL = PatternFill(start_color="1B3A6B", end_color="1B3A6B", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ATTENTION_FILL = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")


def _write_header(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def build_workbook(as_of: date) -> Workbook:
    wb = Workbook()

    dashboard = portfolio_dashboard(PROJECTS, MILESTONES, RISK_LOG, STATUS_REPORTS, as_of)
    attention_ids = {row["project_id"] for row in flag_attention_needed(dashboard)}

    ws_dash = wb.active
    ws_dash.title = "Portfolio Dashboard"
    headers = ["Project", "Name", "Status", "% Complete", "Overdue Milestones", "Open Risks", "Highest Risk Score"]
    _write_header(ws_dash, headers)
    for i, row in enumerate(dashboard, start=2):
        ws_dash.cell(row=i, column=1, value=row["project_id"])
        ws_dash.cell(row=i, column=2, value=row["name"])
        ws_dash.cell(row=i, column=3, value=row["status"])
        ws_dash.cell(row=i, column=4, value=row["percent_complete"])
        ws_dash.cell(row=i, column=5, value=row["overdue_milestone_count"])
        ws_dash.cell(row=i, column=6, value=row["open_risk_count"])
        ws_dash.cell(row=i, column=7, value=row["highest_open_risk_score"])
        if row["project_id"] in attention_ids:
            for col in range(1, len(headers) + 1):
                ws_dash.cell(row=i, column=col).fill = _ATTENTION_FILL

    ws_risk = wb.create_sheet("Risk Log")
    risk_headers = ["Risk ID", "Project", "Description", "Likelihood", "Impact", "Status", "Owner"]
    _write_header(ws_risk, risk_headers)
    for i, risk in enumerate(RISK_LOG, start=2):
        ws_risk.cell(row=i, column=1, value=risk.risk_id)
        ws_risk.cell(row=i, column=2, value=risk.project_id)
        ws_risk.cell(row=i, column=3, value=risk.description)
        ws_risk.cell(row=i, column=4, value=risk.likelihood)
        ws_risk.cell(row=i, column=5, value=risk.impact)
        ws_risk.cell(row=i, column=6, value=risk.status)
        ws_risk.cell(row=i, column=7, value=risk.owner)

    ws_status = wb.create_sheet("Status Reports")
    status_headers = ["Report ID", "Project", "Date", "Summary", "% Complete"]
    _write_header(ws_status, status_headers)
    for i, report in enumerate(STATUS_REPORTS, start=2):
        ws_status.cell(row=i, column=1, value=report.report_id)
        ws_status.cell(row=i, column=2, value=report.project_id)
        ws_status.cell(row=i, column=3, value=report.report_date.isoformat())
        ws_status.cell(row=i, column=4, value=report.summary)
        ws_status.cell(row=i, column=5, value=report.percent_complete)

    for ws in (ws_dash, ws_risk, ws_status):
        for col_cells in ws.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 60)

    return wb


def save_workbook(path: str, as_of: date) -> None:
    build_workbook(as_of).save(path)
