from datetime import date

from openpyxl import load_workbook

from src.report import build_workbook, save_workbook


def test_build_workbook_has_three_sheets():
    wb = build_workbook(date(2026, 6, 20))
    assert wb.sheetnames == ["Portfolio Dashboard", "Risk Log", "Status Reports"]


def test_build_workbook_dashboard_has_header_row():
    wb = build_workbook(date(2026, 6, 20))
    ws = wb["Portfolio Dashboard"]
    header = [c.value for c in ws[1]]
    assert header == ["Project", "Name", "Status", "% Complete", "Overdue Milestones", "Open Risks", "Highest Risk Score"]


def test_build_workbook_dashboard_row_count_matches_projects():
    from src.data_sources import PROJECTS
    wb = build_workbook(date(2026, 6, 20))
    ws = wb["Portfolio Dashboard"]
    assert ws.max_row == len(PROJECTS) + 1  # +1 for header


def test_save_workbook_writes_a_real_readable_file(tmp_path):
    out_path = tmp_path / "test_output.xlsx"
    save_workbook(str(out_path), date(2026, 6, 20))
    assert out_path.exists()

    # Read it back with a fresh load to confirm it's a real, valid xlsx
    # (not just bytes that happen to exist) and the risk log content
    # round-trips correctly.
    wb = load_workbook(str(out_path))
    ws = wb["Risk Log"]
    risk_ids_in_file = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "R01" in risk_ids_in_file
